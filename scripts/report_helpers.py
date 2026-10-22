import os
import sys
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_borders(cell, color="CCCCCC", sz="4", val="single"):
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tcBorders>'
    )
    cell._tc.get_or_add_tcPr().append(borders)

def set_cell_padding(cell, top=140, bottom=140, left=200, right=200):
    mar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    cell._tc.get_or_add_tcPr().append(mar)

def add_heading_styled(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    run = p.add_run(text)
    run.bold = True
    run.font.name = "Times New Roman"
    
    if level == 1:
        p.paragraph_format.space_before = Pt(18)
        p.paragraph_format.space_after = Pt(8)
        run.font.size = Pt(16)
        run.font.color.rgb = RGBColor(15, 44, 89) # Deep Navy
    elif level == 2:
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(30, 64, 175) # Blue
    elif level == 3:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        run.font.size = Pt(13)
        run.font.color.rgb = RGBColor(51, 65, 85) # Slate
    return p

def add_body_paragraph(doc, text, bold_prefix="", italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.bold = True
        r_prefix.font.name = "Times New Roman"
        r_prefix.font.size = Pt(12.5)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.5)
    r.italic = italic
    return p

def add_bullet_item(doc, bold_title, text):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.2
    r_b = p.add_run(bold_title + ": ")
    r_b.bold = True
    r_b.font.name = "Times New Roman"
    r_b.font.size = Pt(12)
    r_t = p.add_run(text)
    r_t.font.name = "Times New Roman"
    r_t.font.size = Pt(12)
    return p

def add_image_with_caption(doc, img_path, caption_title, description):
    if not os.path.exists(img_path):
        print(f"Warning: Image not found: {img_path}")
        return
    
    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img.paragraph_format.space_before = Pt(12)
    p_img.paragraph_format.space_after = Pt(4)
    p_img.paragraph_format.keep_with_next = True
    
    run_img = p_img.add_run()
    run_img.add_picture(str(img_path), width=Inches(6.2))
    
    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_cap.paragraph_format.space_before = Pt(2)
    p_cap.paragraph_format.space_after = Pt(6)
    p_cap.paragraph_format.keep_with_next = True
    
    r_cap_title = p_cap.add_run(caption_title)
    r_cap_title.bold = True
    r_cap_title.font.name = "Times New Roman"
    r_cap_title.font.size = Pt(11)
    r_cap_title.font.color.rgb = RGBColor(15, 23, 42)
    
    if description:
        p_desc = doc.add_paragraph()
        p_desc.paragraph_format.space_before = Pt(0)
        p_desc.paragraph_format.space_after = Pt(12)
        p_desc.paragraph_format.line_spacing = 1.2
        r_desc = p_desc.add_run("↳ Mô tả chức năng: " + description)
        r_desc.italic = True
        r_desc.font.name = "Times New Roman"
        r_desc.font.size = Pt(11)
        r_desc.font.color.rgb = RGBColor(71, 85, 105)

print("Helper definitions loaded successfully.")
